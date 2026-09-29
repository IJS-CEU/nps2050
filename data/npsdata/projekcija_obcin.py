"""Pot občine do leta 2050 po scenariju NPS 2050 (ponazoritev).

1. Razporeditev po razredih (prostor dejanske porabe) za vsak arhetip RenRates v letih 2025, 2030, 2035, 2040 in 2045
   (začetek petletnih obdobij scenarija NPS), za 2050 po vrstah stavb (Class Overview).
2. Prehodi med razredi: monotona sklopitev začetne in končne razporeditve (kot pri Sanki grafu RenRates) – stavbe ohranijo
   vrstni red, slabše stavbe preidejo v razrede, v katere se v scenariju premakne najslabši del fonda.
3. Raba končne energije stavbe v letu y = izhodiščna raba × razmerje intenzivnosti po prehodu; vsota po vrsti stavb se umeri
   na razmerje iz preglednice »Končna raba energije v stavbah v opazovanih letih« (y / 2023).
4. Energenti: izhodiščni deleži občine × državno razmerje energenta po sektorju (preglednica energentov), nato uskladitev
   (IPF) z državnimi vsotami. Emisije: faktorji RenRates po letih, umerjeni na osnutek po sektorju; OVE: en delež OVE v elektriki
   na leto, umerjen na osnutek.
"""
import numpy as np
import openpyxl
import pandas as pd

from .draft import row
from .numbers import num
from .stavbe import CARRIERS, HEAT_CLASSES

HC = HEAT_CLASSES  # A1 … G (najboljši → najslabši)
RR_ORDER = ['G', 'F', 'E', 'D', 'C', 'B2', 'B1', 'A2', 'A1']  # vrstni red v listih RenRates
YEARS = [2030, 2040, 2050]
EF_Y = {'el': [0.22, 0.17, 0.08, 0.03], 'gas': [0.202] * 4, 'elko': [0.267] * 4, 'bio': [0.02] * 4, 'dh': [0.20, 0.17, 0.11, 0.05], 'amb': [0.0] * 4}  # 2025, 2030, 2040, 2050
KWH_GWH = 1e6


def _dist_rr(vals):
    """Vrednosti v vrstnem redu RR_ORDER → vektor v vrstnem redu HC."""
    d = dict(zip(RR_ORDER, vals))
    return np.array([float(d[c] or 0) for c in HC])


def load_paths(ctx) -> dict:
    wb = openpyxl.load_workbook(ctx.files['renrates'], read_only=True, data_only=True)
    dist = {}  # (arhetip, leto) → razporeditev m² po HC
    for sheet, year in (('2026-2030', 2025), ('2031-2035', 2030), ('2036-2040', 2035), ('2041-2045', 2040), ('2046-2050', 2045)):
        for r in list(wb[sheet].iter_rows(values_only=True))[5:]:
            if isinstance(r[0], (int, float)) and r[1] in ('SFH', 'MFH', 'PB', 'SB'):
                dist[(int(r[0]), year)] = _dist_rr(r[3:12])
    co = list(wb['Class Overview'].iter_rows(values_only=True))
    head = co[0][16:25]
    typ2050 = {}
    for r in co[1:12]:
        if isinstance(r[15], str) and r[15].endswith('2050'):
            d = dict(zip(head, r[16:25]))
            typ2050[r[15].split()[0]] = np.array([float(d.get(c) or 0) for c in HC])
    return {'dist': dist, 'typ2050': typ2050}


def _coupling(s, e):
    """Monotona sklopitev (od najslabšega razreda): pogojna verjetnost P(konec j | začetek i)."""
    order = list(range(8, -1, -1))  # G … A1
    S, E = np.cumsum(s[order]), np.cumsum(e[order] * s.sum() / max(e.sum(), 1e-9))
    M = np.zeros((9, 9))
    for ii, i in enumerate(order):
        for jj, j in enumerate(order):
            lo = max(S[ii - 1] if ii else 0, E[jj - 1] if jj else 0)
            M[i, j] = max(0.0, min(S[ii], E[jj]) - lo)
    rs = M.sum(1, keepdims=True)
    return np.where(rs > 0, M / np.maximum(rs, 1e-12), np.eye(9))


def _ratio(Pe, arch, fe_tot, paths, typ_of, years):
    """Razmerje intenzivnosti (leto y / izhodišče) za vsako stavbo."""
    base = np.einsum('nc,nc->n', Pe, fe_tot[arch])
    out = {}
    for y in years:
        r = np.ones(len(arch))
        for a in np.unique(arch):
            m = arch == a
            s = paths['dist'].get((a, 2025))
            if s is None or s.sum() == 0:
                continue
            if y == 2050:
                e45 = paths['dist'][(a, 2045)]
                t = typ_of[a]
                tot45 = sum(paths['dist'][(b, 2045)] for b in typ_of if typ_of[b] == t and (b, 2045) in paths['dist'])
                C45 = _coupling(tot45, paths['typ2050'][t])
                e = e45 @ C45
            else:
                e = paths['dist'][(a, y)]
            C = _coupling(s, e)
            fe_y = (Pe[m] @ C) @ fe_tot[a]
            r[m] = np.where(base[m] > 0, fe_y / base[m], 1.0)
        out[y] = r
    return out


def build(ctx, rr: dict, res_run: dict, nres: pd.DataFrame, kz_si: dict, muni: dict) -> dict:
    paths = load_paths(ctx)
    typ_of = {a: d['typ'] for a, d in rr['arch_all'].items()}
    fe_tot = rr['fe'].sum(2)  # arhetip × razred
    rres = _ratio(res_run['Pe'], res_run['arch'], fe_tot, paths, typ_of, YEARS)
    rnr = _ratio(nres.attrs['Pe'], nres.attrs['arch'], fe_tot, paths, typ_of, YEARS)

    t27 = ctx.draft.table(r'Končna raba energije v stavbah v opazovanih letih')
    path = {k: [num(c) for c in row(t27, p)[2:6]] for k, p in (('hise', r'^Enodružinske'), ('bloki', r'^Večstanovanjske'), ('javne', r'^Stavbe javnega'), ('zasebne', r'^Stavbe zasebnega'))}
    t30 = ctx.draft.table(r'Struktura končne rabe energije po energentih po scenariju NPS 2050')
    pats = {'el': [r'^Električna'], 'amb': [r'^Toplota okolice', r'^Sončna'], 'gas': [r'^Zemeljski plin', r'^Utekočinjeni', r'^Bioplin'], 'elko': [r'^Ekstra lahko'], 'bio': [r'^Lesna'], 'dh': [r'^Daljinska']}
    car_t = {sec: {c: [sum(num(row(t30, p)[off + i]) for p in ps) for i in range(4)] for c, ps in pats.items()} for sec, off in (('res', 2), ('nres', 6))}
    t32 = ctx.draft.table(r'Emisije toplogrednih plinov v stavbah v opazovanih letih')
    tgp_t = {sec: [num(c) * 1e6 for c in row(t32, p)[2:6]] for sec, p in (('res', r'^Stanovanjski'), ('nres', r'^Storitveni'))}
    t33 = ctx.draft.table(r'Obnovljivi viri energije v stavbah v opazovanih letih')
    ove_t = [num(c) / 100 for c in row(t33, r'^Stavbe skupaj')[2:6]]

    res_b = res_run['buildings']
    seg_res = np.where(res_run['cat'] == 'HISA', 'hise', 'bloki')
    ids = sorted(muni)
    sdh = kz_si['sdh']  # delež OVE daljinske toplote po občini (Series)
    result = {e: {} for e in ids}
    rres[2023] = np.ones(len(res_run['arch'])); rnr[2023] = np.ones(len(nres))
    for yi, y in enumerate([2023] + YEARS):
        fe = []
        for b, r, seg in ((res_b, rres[y], seg_res), (nres, rnr[y], nres.seg.to_numpy())):
            f = b.fe.to_numpy() * r
            for s_ in np.unique(seg):
                m = seg == s_
                tgt = b.fe.to_numpy()[m].sum() * path[s_][yi] / path[s_][0]
                f[m] *= tgt / f[m].sum()
            fe.append(pd.Series(f, index=b.obcina.to_numpy()).groupby(level=0).sum().reindex(ids, fill_value=0))
        agg = {}
        for sec, b, fy in (('res', res_b, fe[0]), ('nres', nres, fe[1])):
            base = b.groupby('obcina')[[f'c_{c}' for c in CARRIERS]].sum().reindex(ids, fill_value=0).to_numpy()
            sh = base / np.maximum(base.sum(1, keepdims=True), 1e-9)
            k = np.array([car_t[sec][c][yi] / max(car_t[sec][c][0], 1e-9) for c in CARRIERS])
            X = sh * k
            X = X / np.maximum(X.sum(1, keepdims=True), 1e-9) * fy.to_numpy()[:, None]
            tgt = np.array([car_t[sec][c][yi] for c in CARRIERS]); tgt = tgt / tgt.sum() * X.sum()
            for _ in range(200):  # IPF: vsote občin (vrstice) in državne vsote energentov (stolpci)
                X *= np.where(X.sum(0) > 0, tgt / np.maximum(X.sum(0), 1e-9), 1)
                X *= (fy.to_numpy() / np.maximum(X.sum(1), 1e-9))[:, None]
            agg[sec] = pd.DataFrame(X, index=ids, columns=CARRIERS)
        # emisije in OVE
        ef_dh = EF_Y['dh'][yi] * (1 - sdh) / (1 - kz_si['s_bar'])
        tg = {}
        for sec, A in agg.items():
            raw = sum(A[c] * EF_Y[c][yi] for c in CARRIERS if c != 'dh') + A['dh'] * ef_dh
            tg[sec] = raw * tgp_t[sec][yi] / raw.sum()
        fe_all = sum(A.sum(1) for A in agg.values())
        rd = sum(A.bio + A.amb + A.dh * sdh for A in agg.values())
        el = sum(A.el for A in agg.values())
        s_el = (ove_t[yi] * fe_all.sum() - rd.sum()) / el.sum()
        ctx.check(0 <= s_el <= 1, f'projekcija {y}: delež OVE v elektriki {100 * s_el:.1f} %')
        ove = rd + el * s_el
        for e in ids:
            result[e][str(y)] = {'fe_gwh': round(float(fe_all[e]) / KWH_GWH, 1), 'fe_res_gwh': round(float(agg['res'].loc[e].sum()) / KWH_GWH, 1),
                                 'tgp_kt': round(float(tg['res'][e] + tg['nres'][e]) / 1e6, 2),
                                 'ove_pct': round(100 * float(ove[e]) / float(fe_all[e]), 1) if fe_all[e] else None,
                                 'fossil_pct': round(100 * float(sum(agg[s_].loc[e, ['gas', 'elko']].sum() for s_ in agg)) / float(fe_all[e]), 1) if fe_all[e] else None}
        ctx.check(abs(sum(v[str(y)]['tgp_kt'] for v in result.values()) - (tgp_t['res'][yi] + tgp_t['nres'][yi]) / 1e6) < 2, f'projekcija {y}: emisije = osnutek')
    return {str(muni[e][0]): v for e, v in result.items()}
