"""11.3 Drsnik »Kaj pa, če«: kje pristane trajektorija stanovanjskih stavb pri drugačni stopnji prenove.

Vir: pet zagonov modela RenRates_SI (RenRates_SI_drsnik_k070 … k200.xlsx ob RenRates_SI_NPS_scenariji.xlsx), ki se razlikujejo
samo v stopnjah prenove enostanovanjskih in večstanovanjskih stavb (3_Rates F6:J7). Ime k… je faktor glede na scenarij S2 pred
5. 10. 2026; od takrat je scenarij NPS (S2) zagon k125 (stopnje hiš in blokov × 1,25, da doseže mejnik EPBD 2030 v celotnem razponu
ocene stanja 2025), faktor v podatkih pa je glede na novi S2.
Krivulje se ne interpolirajo in ne izmišljujejo (CLAUDE.md §11.3).

- Povprečna primarna energija stanovanjskih stavb pri fiksnih faktorjih 2025: energenti SFH + MFH (8_SFH_Charts) × faktorji PE
  2025 (7_Carriers) / stanovanjska površina; kot indeks (2025 = 100).
- kWh/(m²·a) na lestvici trajektorije NPS: ocenjeno stanje 2025 × indeks položaja (2025 = 100). Model se začne leta 2025, zato je
  stanje 2025 ocenjeno iz podatkov (OCENA_2025), ne privzeto s premice 2020–2030: (a) registri 2020–2025 (kataster in Eko sklad,
  ukrepi ovrednoteni z arhetipi iz izkaznic, z novogradnjami): 254–256; (b) primarna raba stanovanjskih stavb 2020 → 2023 iz
  izhodiščne preglednice osnutka (NEPN 2024), popravljena za stopinjske dneve (Eurostat) in rast površine (NEPN 2020/2024),
  2023–2025 s tempom sedanje prakse (S0): okoli 257. Razpon 252–260.
- Dosežena stopnja prenove: prenovljena stanovanjska površina 2026–2050 (delna, celovita, celovita s potresno) / površina / 25 let.
  Model prenavlja le stavbe razredov C–G, zato dosežena stopnja zaostaja za vpisano in se pri visokih stopnjah nasiči.
"""
import json

import openpyxl

from .context import OUT, Context, write_csv, write_json

POS = [('k070', 0.56), ('k100', 0.8), ('k125', 1.0), ('k150', 1.2), ('k200', 1.6)]
OCENA_2025 = {'osrednja': 255, 'spodnja': 252, 'zgornja': 260}  # kWh/(m²·a), glej opis zgoraj
VIRI_2025 = {'registri': [254, 256], 'bilanca': 257}  # data/raw/ocena_2020_2025.py (okni 2020–2025 in 2021–2025); bilanca s popravkom za vreme
BASE_2020 = 269
PERIODS = ['2026-2030', '2031-2035', '2036-2040', '2041-2045', '2046-2050']
YEARS = [2025, 2030, 2035, 2040, 2045, 2050]
CARRIER_ROWS = {'SFH': 147, 'MFH': 231}
# Enotni stroški in delež javnih sredstev P / D / DS (Scenario_Compare vrstici 142–143, predpostavka NPS 7.1)
COST = {'SFH': (200, 520, 750), 'MFH': (170, 420, 620), 'PB': (280, 700, 1000), 'SB': (250, 650, 950)}
PUBLIC = {'SFH': (.15, .30, .40), 'MFH': (.15, .30, .40), 'PB': (.5, .5, .5), 'SB': (.10, .20, .25)}


def _run(path, s0=False) -> dict:
    """s0: v S0 je strošek delne (lažje) prenove 120 €/m² (hiše) in 100 €/m² (bloki), Scenario_Compare vrstica 85."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rates = wb['3_Rates']
    rr = {i: r for i, r in enumerate(rates.iter_rows(min_row=1, max_row=70, max_col=10, values_only=True), start=1)}
    typ = {int(rr[i][0]): rr[i][3] for i in range(41, 71)}
    stock = {}
    for i in range(41, 71):
        stock[rr[i][3]] = stock.get(rr[i][3], 0) + (rr[i][2] or 0)
    res_area = stock['SFH'] + stock['MFH']
    ren = inv = inv30 = pub = 0.0
    per_rate = {t: [[0.0, 0.0] for _ in PERIODS] for t in ('SFH', 'MFH')}  # [vse, celovito] m² po obdobjih
    for pi, per in enumerate(PERIODS):
        ws = wb[per]
        rows = {i: r for i, r in enumerate(ws.iter_rows(min_row=1, max_row=110, max_col=32, values_only=True), start=1)}
        for k, (rng, col) in enumerate(((range(6, 36), 31), (range(42, 72), 28), (range(78, 108), 28))):  # delna AF, celovita AC, celovita + potresna AC
            for i in rng:
                t = typ[int(rows[i][0])]
                a = rows[i][col] or 0
                if t in ('SFH', 'MFH'):
                    ren += a
                    per_rate[t][pi][0] += a
                    if k:
                        per_rate[t][pi][1] += a
                c = a * ({'SFH': 120, 'MFH': 100}[t] if s0 and k == 0 and t in ('SFH', 'MFH') else COST[t][k]) / 1e9
                inv += c; pub += c * PUBLIC[t][k]
                if pi == 0:
                    inv30 += c
    ch = {i: r for i, r in enumerate(wb['8_SFH_Charts'].iter_rows(min_row=1, max_row=240, max_col=7, values_only=True), start=1)}
    car = wb['7_Carriers']
    pef = [car.cell(22 + j, 2).value for j in range(5)]
    pe = [sum(ch[CARRIER_ROWS[t] + j][1 + y] * pef[j] for t in CARRIER_ROWS for j in range(5)) * 1e6 / res_area for y in range(6)]
    co = {i: r for i, r in enumerate(wb['Class Overview'].iter_rows(min_row=1, max_row=156, max_col=7, values_only=True), start=1)}
    sc = {i: r for i, r in enumerate(wb['Scenario_Compare'].iter_rows(min_row=1, max_row=43, max_col=2, values_only=True), start=1)}
    lab = {str(r[0]).strip(): r[1] for r in sc.values() if r[0]}
    return {
        'input_pct': {'SFH': [100 * v for v in rr[6][5:10]], 'MFH': [100 * v for v in rr[7][5:10]]},
        'rate_res_pct': 100 * ren / res_area / 25,
        # stopnja po obdobjih [% površine na leto]: skupaj (hiše + bloki, uteženo s površino) in ločeno
        'per_pct': {
            'skupaj': [[100 * (per_rate['SFH'][i][j] + per_rate['MFH'][i][j]) / res_area / 5 for j in (0, 1)] for i in range(5)],
            'SFH': [[100 * per_rate['SFH'][i][j] / stock['SFH'] / 5 for j in (0, 1)] for i in range(5)],
            'MFH': [[100 * per_rate['MFH'][i][j] / stock['MFH'] / 5 for j in (0, 1)] for i in range(5)],
        },
        'pe_fixed': pe,
        'fe_gwh': list(co[87][1:7]),
        'fe_2050_live': sc[14][1],
        'fg_pct': [100 * v for v in co[151][1:7]],
        'epbd_red': {'2030': co[142][2], '2035': co[143][2]},
        'red_2050': {'fe': lab['Final energy reduction vs 2025'], 'pe': lab['Primary energy reduction vs 2025'], 'co2': lab['CO₂ reduction vs 2025']},
        'inv_bn': inv, 'inv_2030_bn': inv30, 'public_bn': pub,
    }


def build(ctx: Context) -> dict:
    base = ctx.files['renrates']
    runs = {}
    for k, f in POS:
        p = base.with_name(f'RenRates_SI_drsnik_{k}.xlsx')
        ctx.check(p.exists(), f'drsnik: datoteka {p.name}')
        runs[k] = _run(p)
    # k125 je scenarij S2: mora se ujemati z živim modelom v RenRates_SI_NPS_scenariji.xlsx
    s2 = openpyxl.load_workbook(base, read_only=True, data_only=True)['Scenario_Compare']
    fe50_s2 = next(r[1] for r in s2.iter_rows(min_row=12, max_row=20, max_col=2, values_only=True) if r[0] == 'Final energy 2050 [GWh/yr]')
    ctx.check(abs(runs['k125']['fe_2050_live'] - fe50_s2) < 0.01, f'drsnik: k125 = S2 (končna energija 2050 {runs["k125"]["fe_2050_live"]:.1f} = {fe50_s2:.1f} GWh)')
    ctx.check(all(runs[a]['fe_gwh'][5] > runs[b]['fe_gwh'][5] for (a, _), (b, _) in zip(POS, POS[1:])), 'drsnik: višja stopnja → nižja končna energija 2050')

    tr = json.loads((OUT / 'trajektorija.json').read_text(encoding='utf-8'))
    k = tr['kwh_m2']
    nps = dict(zip(k['years'], k['nps']))
    nps[2025] = round((nps[2020] + nps[2030]) / 2)
    epbd = {str(e['year']): e for e in k['epbd_max']}
    idx = {key: [100 * v / r['pe_fixed'][0] for v in r['pe_fixed']] for key, r in runs.items()}
    out = []
    for key, f in POS:
        r = runs[key]
        kwh = [nps[2020]] + [round(OCENA_2025['osrednja'] * i / 100) for i in idx[key]]
        lo = [nps[2020]] + [round(OCENA_2025['spodnja'] * i / 100) for i in idx[key]]
        hi = [nps[2020]] + [round(OCENA_2025['zgornja'] * i / 100) for i in idx[key]]
        out.append({
            'id': key, 'faktor': f,
            'stopnja_dosezena_pct': round(r['rate_res_pct'], 2),
            'stopnja_vpisana_2026_2030_pct': {t: round(v[0], 2) for t, v in r['input_pct'].items()},
            'kwh_m2': kwh, 'kwh_m2_spodnja': lo, 'kwh_m2_zgornja': hi,
            'zmanjsanje_2030_pct': round(100 * (1 - kwh[2] / BASE_2020), 1), 'zmanjsanje_2035_pct': round(100 * (1 - kwh[3] / BASE_2020), 1),
            'indeks_pe': [round(v, 1) for v in idx[key]],
            'fe_twh': [round(v / 1000, 2) for v in r['fe_gwh']],
            'fe_zmanjsanje_2050_pct': round(100 * (1 - r['fe_gwh'][5] / r['fe_gwh'][0])),
            'fg_pct': [round(v, 1) for v in r['fg_pct']],
            'epbd_2030': kwh[2] <= epbd['2030']['max'], 'epbd_2035': kwh[3] <= epbd['2035']['max'],
            'zmanjsanje_2050_pct': {k: round(100 * v, 1) for k, v in r['red_2050'].items()},
            'nalozbe_mrd': round(r['inv_bn'], 1), 'nalozbe_2026_2030_mrd': round(r['inv_2030_bn'], 1), 'javna_mrd': round(r['public_bn'], 2),
        })
    # Referenca: S0 – nadaljevanje sedanje prakse (zagon modela po definiciji v Scenario_Compare, enaka kalibracija)
    p0 = base.with_name('RenRates_SI_drsnik_S0.xlsx')
    ctx.check(p0.exists(), f'drsnik: datoteka {p0.name}')
    runs['S0'] = _run(p0, s0=True)
    idx['S0'] = [100 * v / runs['S0']['pe_fixed'][0] for v in runs['S0']['pe_fixed']]
    ctx.check(abs(runs['S0']['pe_fixed'][0] - runs['k125']['pe_fixed'][0]) < 1e-6, 'drsnik: S0 in položaji imajo isto izhodišče 2025')
    ref = None
    for key in ['S0'] + [k for k, _ in POS]:
        r = runs[key]
        ps = r['per_pct']
        extra = {'stopnja_2026_2030_pct': round(ps['skupaj'][0][0], 2), 'celovito_2026_2030_pct': round(ps['skupaj'][0][1], 2),
                 'stopnje_obdobja_pct': [round(v[0], 2) for v in ps['skupaj']], 'celovito_obdobja_pct': [round(v[1], 2) for v in ps['skupaj']],
                 'hise_2026_2030_pct': [round(v, 2) for v in ps['SFH'][0]], 'bloki_2026_2030_pct': [round(v, 2) for v in ps['MFH'][0]]}
        if key == 'S0':
            kwh = [nps[2020]] + [round(OCENA_2025['osrednja'] * i / 100) for i in idx['S0']]
            ref = {'id': 'S0', 'ime': 'sedanja praksa (S0)', 'stopnja_dosezena_pct': round(r['rate_res_pct'], 2), 'kwh_m2': kwh,
                   'kwh_m2_spodnja': [nps[2020]] + [round(OCENA_2025['spodnja'] * i / 100) for i in idx['S0']],
                   'kwh_m2_zgornja': [nps[2020]] + [round(OCENA_2025['zgornja'] * i / 100) for i in idx['S0']],
                   'zmanjsanje_2030_pct': round(100 * (1 - kwh[2] / BASE_2020), 1), 'zmanjsanje_2035_pct': round(100 * (1 - kwh[3] / BASE_2020), 1),
                   'fg_pct': [round(v, 1) for v in r['fg_pct']], 'epbd_2030': kwh[2] <= epbd['2030']['max'], 'epbd_2035': kwh[3] <= epbd['2035']['max'],
                   'zmanjsanje_2050_pct': {k2: round(100 * v, 1) for k2, v in r['red_2050'].items()},
                   'nalozbe_mrd': round(r['inv_bn'], 1), 'nalozbe_2026_2030_mrd': round(r['inv_2030_bn'], 1), 'javna_mrd': round(r['public_bn'], 2), **extra}
        else:
            next(o for o in out if o['id'] == key).update(extra)
    ctx.check(ref['stopnja_2026_2030_pct'] > 1.5 and ref['celovito_2026_2030_pct'] < ref['stopnja_2026_2030_pct'] / 3,
              f"drsnik: S0 {ref['stopnja_2026_2030_pct']} % na leto, od tega celovito {ref['celovito_2026_2030_pct']} %")
    ctx.check(nps[2020] == BASE_2020, f'drsnik: izhodišče 2020 = {BASE_2020} kWh/(m²·a)')
    data = {
        'meta': ctx.meta(['Model RenRates_SI (uskladitev z NEPN 2024, 3. 10. 2026; scenarij NPS s stopnjami × 1,25, 5. 10. 2026): pet zagonov z različnimi stopnjami prenove stanovanjskih stavb',
                          f'{ctx.draft.name}: nacionalna trajektorija stanovanjskega fonda in najvišje vrednosti po EPBD'],
                         note='kWh/(m²·a): ocenjeno stanje 2025 (registri 2020–2025 in energetska bilanca s popravkom za vreme; 255, razpon 252–260) × indeks položaja iz modela. '
                              'Spreminjajo se samo stopnje prenove enostanovanjskih in večstanovanjskih stavb; vse drugo je kot v scenariju NPS 2050.'),
        'years': [2020] + YEARS, 'nps': [nps[2020]] + [None if y == 2025 else nps[y] for y in YEARS], 'epbd_max': k['epbd_max'],
        'privzeto': 'k125', 'polozaji': out, 'referenca': ref, 'ocena_2025': OCENA_2025, 'viri_2025': VIRI_2025,
    }
    write_json('drsnik', data)
    write_csv('drsnik', ['položaj', 'stopnja prenove stanovanjskih stavb 2026–2030 [%/leto]', 'od tega celovito [%/leto]', 'hiše 2026–2030 [%/leto]', 'bloki 2026–2030 [%/leto]',
                         'povprečna stopnja 2026–2050 [%/leto]'] + [f'{y} [kWh/(m²·a)]' for y in data['years']] + ['zmanjšanje končne energije 2025–2050 [%]'],
              [[p['id'], p['stopnja_2026_2030_pct'], p['celovito_2026_2030_pct'], p['hise_2026_2030_pct'][0], p['bloki_2026_2030_pct'][0], p['stopnja_dosezena_pct'],
                *p['kwh_m2'], p['zmanjsanje_2050_pct']['fe']] for p in [ref] + out])
    return data
