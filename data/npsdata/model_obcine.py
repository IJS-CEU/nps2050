"""Modelska ocena rabe energije stanovanjskih stavb po občinah (CLAUDE.md §10.1, razširitev).

Postopek (podrobno v docs/ODLOCITVE.md in na strani »Moja občina – metoda«):
1. Vsaka stanovanjska stavba iz katastra dobi arhetip RenRates (vrsta × obdobje gradnje).
2. Računski prostor (razredi toplotnih potreb A1–G, kot v izkaznicah): izhodišče je osnovna (nekalibrirana) razporeditev
   arhetipa iz RenRates (1_Calib). Stavbe z računsko izkaznico imajo znan razred. Pri drugih se razporeditev nagne glede
   na znane ukrepe (Eko sklad, vpisane obnove v katastru) – moč nagiba je ocenjena na stavbah z izkaznico.
3. Uskladitev (raking): vsota po arhetipu in razredu se ujema z RenRates; po občinah se potreba po toploti nagne proti
   toplotni karti 2020 (omejeno), državna vsota ostane nespremenjena.
4. Prostor dejanske porabe: vsaka stavba se prenese z istim operatorjem λ kot v kalibraciji RenRates (1_Calib), zato se
   država točno ujema s kalibriranim RenRates (bilanca 2024, 33,54 PJ).
5. Energenti: izkaznica ali Eko sklad, kjer sta znana; drugje mešanica arhetipa. Državne vsote po energentih se uskladijo
   s preglednico 30 osnutka (stanovanjske stavbe 2023).
6. Razredi NPS 2050 in delež nad pragom 43 %: iz računskega prostora prek preslikave razred toplotnih potreb → razred NPS,
   izmerjene na stavbah z izkaznico.
Objavijo se samo seštevki po občinah z razponom (različice uskladitve s toplotno karto).
"""
import json
import re
import sqlite3
import struct

import numpy as np
import openpyxl
import pandas as pd
import shapely
from shapely.geometry import shape

from .context import OUT, Context, write_csv, write_json
from .draft import row
from .numbers import num
from . import cilji_obcin, kazalniki_obcin, model_nres, projekcija_obcin, skupine_obcin
from .context import DATA_DIR
from .meritve import build_measured
from .stavbe import CARRIERS, HEAT_CLASSES, build_table

HC = HEAT_CLASSES
Q_MID = np.array([5, 12.5, 20, 30, 47.5, 82.5, 127.5, 180, 250.0])  # sredine razredov toplotnih potreb, kWh/(m²·a)
NPS = ['A', 'B', 'C', 'D', 'E', 'F', 'G']
BOUNDS = ['A|B', 'B|C', 'C|D', 'D|E', 'E|F', 'F|G']
# Faktorji primarne energije za preglednico 30 osnutka (izhodiščno leto).
FP = {'el': 2.5, 'amb': 1.0, 'gas': 1.1, 'elko': 1.1, 'bio': 1.2, 'dh': 1.23}
KWH_PER_KTOE = 11.63e6
# Različice uskladitve s toplotno karto (spodnja in zgornja meja faktorja občine); osrednja je druga.
EI_WEIGHT = 0.85
EI_SHRINK = 30  # navidezne stavbe državnega povprečja pri krajevni mešanici energentov
RENOV_FROM = 2010
VARIANTS = {'brez': (1.0, 1.0), 'osrednja': (0.8, 1.25), 'mocna': (0.6, 1.6)}


# ---------------------------------------------------------------- RenRates
def load_renrates(ctx: Context) -> dict:
    wb = openpyxl.load_workbook(ctx.files['renrates'], read_only=True, data_only=True)
    cal = list(wb['1_Calib'].iter_rows(values_only=True))
    inp = list(wb['1_Inputs'].iter_rows(values_only=True))
    lam = float(cal[4][1])
    lam_s = float(cal[5][1])
    arch, arch_all = {}, {}
    for i in range(30):
        r = inp[10 + i]
        arch_all[int(r[0])] = {'name': r[1], 'typ': r[2], 'stock': float(r[3] or 0), 'base': np.array(cal[9 + i][4:13], float), 'final': np.array(r[4:13], float)}
        if r[2] not in ('SFH', 'MFH'):
            continue
        arch[int(r[0])] = {'name': r[1], 'typ': r[2], 'stock': float(r[3] or 0),
                           'base': np.array(cal[9 + i][4:13], float), 'final': np.array(r[4:13], float)}
    fe = np.zeros((31, 9, 5))  # arhetip × razred × (el, gas, elko, bio, dh) kWh/m² ogrevane površine
    for r in list(wb['5_FinalEnergy'].iter_rows(values_only=True))[3:]:
        if r[0] is None or r[2] not in HC:
            continue
        fe[int(r[0]), HC.index(r[2])] = [r[18] or 0, r[19] or 0, r[20] or 0, r[21] or 0, r[22] or 0]
    T = np.zeros((9, 9))
    T[0, 0] = 1
    for j in range(1, 9):
        T[j, j], T[j, j - 1] = 1 - lam, lam
    T3 = np.linalg.matrix_power(T, 3)
    Ts = np.zeros((9, 9))
    Ts[0, 0] = 1
    for j in range(1, 9):
        Ts[j, j], Ts[j, j - 1] = 1 - lam_s, lam_s
    T3s = np.linalg.matrix_power(Ts, 3)
    for a, d in arch.items():
        ctx.check(np.abs(d['base'] @ T3 - d['final']).max() < 1e-6, f'model: operator λ reproducira kalibracijo RenRates ({d["name"]})')
    ok_s = all(np.abs(d['base'] @ T3s - d['final']).max() < 1e-6 for d in arch_all.values() if d['typ'] not in ('SFH', 'MFH') and d['stock'] > 0)
    ctx.check(ok_s, 'model: operator λ za storitve reproducira kalibracijo RenRates (vsi nestanovanjski arhetipi)')
    return {'lam': lam, 'arch': arch, 'arch_all': arch_all, 'fe': fe, 'T3': T3, 'T3s': T3s}


def archetype(df: pd.DataFrame) -> pd.Series:
    y = df.year.fillna(1970)
    sfh = np.select([y <= 1920, y <= 1965, y <= 1981, y <= 2008], [1, 2, 3, 4], 5)
    mfh = np.select([y <= 1965, y <= 1981, y <= 2008], [6, 8, 10], 11)
    return pd.Series(np.where(df.cat == 'HISA', sfh, mfh), index=df.index)


# ---------------------------------------------------------------- toplotna karta
def _tm_inverse(E, N):
    """D96/TM (EPSG:3794) → geografske koordinate ETRS89 (Snyder, prečna Mercatorjeva, GRS80)."""
    a, f = 6378137.0, 1 / 298.257222101
    e2 = f * (2 - f); ep2 = e2 / (1 - e2); k0 = 0.9999; lon0 = np.radians(15)
    x, y = E - 500000.0, N + 5000000.0
    M = y / k0
    mu = M / (a * (1 - e2 / 4 - 3 * e2 ** 2 / 64 - 5 * e2 ** 3 / 256))
    e1 = (1 - np.sqrt(1 - e2)) / (1 + np.sqrt(1 - e2))
    phi1 = (mu + (3 * e1 / 2 - 27 * e1 ** 3 / 32) * np.sin(2 * mu) + (21 * e1 ** 2 / 16 - 55 * e1 ** 4 / 32) * np.sin(4 * mu)
            + (151 * e1 ** 3 / 96) * np.sin(6 * mu) + (1097 * e1 ** 4 / 512) * np.sin(8 * mu))
    C1 = ep2 * np.cos(phi1) ** 2; T1 = np.tan(phi1) ** 2
    N1 = a / np.sqrt(1 - e2 * np.sin(phi1) ** 2); R1 = a * (1 - e2) / (1 - e2 * np.sin(phi1) ** 2) ** 1.5
    D = x / (N1 * k0)
    lat = phi1 - (N1 * np.tan(phi1) / R1) * (D ** 2 / 2 - (5 + 3 * T1 + 10 * C1 - 4 * C1 ** 2 - 9 * ep2) * D ** 4 / 24
                                             + (61 + 90 * T1 + 298 * C1 + 45 * T1 ** 2 - 252 * ep2 - 3 * C1 ** 2) * D ** 6 / 720)
    lon = lon0 + (D - (1 + 2 * T1 + C1) * D ** 3 / 6 + (5 - 2 * C1 + 28 * T1 - 3 * C1 ** 2 + 8 * ep2 + 24 * T1 ** 2) * D ** 5 / 120) / np.cos(phi1)
    return np.degrees(lon), np.degrees(lat)


def heat_map_by_municipality(ctx: Context) -> pd.Series:
    """Potreba po toploti za ogrevanje (vse stavbe) iz toplotne karte 2020, seštevek po občinah (MWh)."""
    con = sqlite3.connect(ctx.files['toplotna_karta'])
    rows = con.execute('select geom, QnH_og_MWh from TK_og2020_EPSG3794 where QnH_og_MWh > 0').fetchall()
    xs, ys, q = [], [], []
    for g, v in rows:
        flags = g[3]
        env = (flags >> 1) & 7
        hlen = 8 + {0: 0, 1: 32, 2: 48, 3: 48, 4: 64}[env]
        c = shapely.from_wkb(bytes(g[hlen:])).centroid
        xs.append(c.x); ys.append(c.y); q.append(v)
    lon, lat = _tm_inverse(np.array(xs), np.array(ys))
    ob = json.loads(ctx.files['obcine'].read_text(encoding='utf-8'))
    polys = [shape(f['geometry']) for f in ob['features']]
    ids = [str(f['properties']['EID_OBCINA']) for f in ob['features']]
    tree = shapely.STRtree(polys)
    pts = shapely.points(lon, lat)
    pi, gi = tree.query(pts, predicate='within')
    s = pd.Series(np.array(q)[pi], index=[ids[i] for i in gi]).groupby(level=0).sum()
    ctx.check(s.sum() / sum(q) > 0.995, f'model: toplotna karta – {100 * s.sum() / sum(q):.1f} % potrebe pripisane občinam')
    return s


# ---------------------------------------------------------------- pomožno
def _rake_rows(P, w, target, fixed, iters=500):
    """IPF: vrstice (stavbe) so porazdelitve (vsota 1), stolpčne vsote Σ w·P se prilagodijo cilju; fiksne vrstice ostanejo."""
    free = ~fixed
    tgt = target - (w[fixed, None] * P[fixed]).sum(0)
    tgt = np.maximum(tgt, 1e-9 * target.sum())
    for _ in range(iters):
        cur = (w[free, None] * P[free]).sum(0)
        k = np.where(cur > 0, tgt / np.maximum(cur, 1e-12), 1.0)
        if np.abs(k - 1).max() < 1e-7:
            break
        P[free] *= k
        P[free] /= P[free].sum(1, keepdims=True)
    return P


def _tilt(P, q, beta):
    X = P * np.exp(beta[:, None] * (q[None, :] - q.mean()) / 100)
    return X / X.sum(1, keepdims=True)


def run(ctx: Context, df: pd.DataFrame, rr: dict, fond: dict, tk: pd.Series | None, clip: tuple[float, float]) -> dict:
    """Ena različica modela. Vrne seštevke po občinah in državi."""
    res = df[df.cat.isin(['HISA', 'BLOKI'])].copy()
    res['arch'] = archetype(res)
    n = len(res)
    # uteži: ogrevana površina RenRates, razdeljena po uporabni površini katastra znotraj arhetipa
    area_a = res.groupby('arch').m2.sum()
    stock = {a: d['stock'] for a, d in rr['arch'].items()}
    w = (res.m2 * res.arch.map(lambda a: stock[a] / area_a[a])).to_numpy()
    arch = res.arch.to_numpy()

    # 1–2: izhodišče in dokazi v računskem prostoru
    y = lambda c: pd.to_numeric(res[c], errors='coerce')
    # Obnove iz katastra štejejo kot energetske le od leta 2010 (prej vpisi pogosto pomenijo le barvanje fasade,
    # menjavo kritine ipd.; odločitev Gašperja 29. 9. 2026).
    env = (res.es_ovoj.notna() | (y('obnova_fasada') >= RENOV_FROM) | (y('obnova_streha') >= RENOV_FROM)).to_numpy()
    win = (res.es_okna.notna() | (y('obnova_okna') >= RENOV_FROM)).to_numpy()
    sysn = (res.es_tc.notna() | res.es_biomasa.notna() | (y('obnova_instal') >= RENOV_FROM)).to_numpy()
    flag = env * 4 + win * 2 + sysn
    known = res.ei_razred.notna().to_numpy()
    kc = res.ei_razred.map({c: i for i, c in enumerate(HC)}).fillna(-1).astype(int).to_numpy()
    obc = res.obcina.to_numpy()
    # Skupine stavb z enako porazdelitvijo med uskladitvijo: (arhetip, občina, zastavice, razred izkaznice).
    keys = pd.DataFrame({'a': arch, 'o': obc, 'f': flag, 'k': kc})
    gid, uniq = pd.factorize(pd.MultiIndex.from_frame(keys))
    ug = pd.DataFrame(list(uniq), columns=["a", "o", "f", "k"])
    ng = len(ug)
    wg = np.bincount(gid, weights=w, minlength=ng)
    ga, go, gf, gk = ug.a.to_numpy(), ug.o.to_numpy(), ug.f.to_numpy(), ug.k.to_numpy()
    gknown = gk >= 0
    # razmerje verjetnosti razreda glede na zastavice, ocenjeno na stavbah z izkaznico (po vrsti stavbe, glajenje +1)
    gtyp = np.where(ga <= 5, 'HISA', 'BLOKI')
    P = np.stack([rr['arch'][a]['base'] for a in ga])
    for typ in ('HISA', 'BLOKI'):
        m_typ = (res.cat == typ).to_numpy()
        ref = np.bincount(kc[m_typ & known], minlength=9) + 1.0
        ref /= ref.sum()
        for f in range(8):
            h = np.bincount(kc[m_typ & known & (flag == f)], minlength=9) + 1.0
            P[(gtyp == typ) & ~gknown & (gf == f)] *= (h / h.sum()) / ref
    P /= P.sum(1, keepdims=True)
    # Izkaznica je močan, a ne absoluten dokaz (85 % na razredu iz izkaznice), da se vsote po arhetipu vedno ujemajo z RenRates.
    P[gknown] = EI_WEIGHT * np.eye(9)[gk[gknown]] + (1 - EI_WEIGHT) * P[gknown]
    nofix = np.zeros(ng, bool)

    # 3: uskladitev po arhetipih in nagib po občinah proti toplotni karti
    targets = {a: rr['arch'][a]['stock'] * rr['arch'][a]['base'] for a in rr['arch']}
    nres = df[~df.cat.isin(['HISA', 'BLOKI'])]
    rake = lambda P: [P.__setitem__(ga == a, _rake_rows(P[ga == a], wg[ga == a], targets[a], nofix[ga == a])) for a in rr['arch']]
    rake(P)
    factor = pd.Series(1.0, index=np.unique(obc))
    factor0 = factor
    if tk is not None and clip != (1.0, 1.0):
        idx = pd.Series(np.arange(ng)).groupby(go).apply(lambda s: s.to_numpy())
        for it in range(12):
            H = pd.Series(wg * (P @ Q_MID)).groupby(go).sum()
            q_res = H.sum() / res.m2.sum()
            Hn = nres.groupby('obcina').m2.sum().reindex(H.index, fill_value=0) * q_res
            tkm = tk.reindex(H.index, fill_value=0) * 1000  # kWh
            f = (tkm / tkm.sum()) / ((H + Hn) / (H + Hn).sum())
            if it == 0:
                factor0 = f  # izhodiščno razmerje toplotna karta / model (za prikaz)
            factor = f.clip(*clip)
            for m_id, fac in factor.items():
                ii = idx[m_id]
                fr = ii
                if abs(fac - 1) < 1e-3:
                    continue
                Pm, wm = P[fr], wg[fr]
                goal = H[m_id] * fac
                lo, hi = -20.0, 20.0
                for _ in range(40):
                    mid = (lo + hi) / 2
                    val = (wm * (_tilt(Pm, Q_MID, np.full(len(fr), mid)) @ Q_MID)).sum()
                    lo, hi = (mid, hi) if val < goal else (lo, mid)
                P[fr] = _tilt(Pm, Q_MID, np.full(len(fr), (lo + hi) / 2))
            rake(P)
    P = P[gid]
    known = gknown[gid]
    kc = np.where(known, kc, 0)

    # 4: prostor dejanske porabe
    Pe = P @ rr['T3']
    fe = rr['fe'][arch]  # n × 9 × 5
    FEc = np.einsum('nc,nck->nk', Pe, fe) * w[:, None]  # kWh po (el, gas, elko, bio, dh)
    tot = FEc.sum(1)

    # 5: energenti (el, amb, gas, elko, bio, dh)
    C = np.zeros((n, 6))
    C[:, 0], C[:, 2], C[:, 3], C[:, 4], C[:, 5] = FEc[:, 0], FEc[:, 1], FEc[:, 2], FEc[:, 3], FEc[:, 4]
    C[:, 1] = 0.05 * tot
    C[:, [0, 2, 3, 4, 5]] *= 0.95
    # Daljinska toplota le v občinah, kjer jo potrjujejo izkaznice (vsaj 10 stavb z vsaj 50 % daljinske toplote);
    # drugje se njen delež razdeli med ostala goriva arhetipa.
    dh_m = set(df[df.ei_sh_dh.fillna(0) >= 0.5].groupby('obcina').size().loc[lambda x: x >= 10].index)
    no_dh = ~np.isin(obc, list(dh_m))
    other = C[no_dh][:, [2, 3, 4]]
    osum = other.sum(1, keepdims=True)
    share = np.where(osum > 0, other / np.maximum(osum, 1e-9), np.array([[0.2, 0.3, 0.5]]))
    C[np.ix_(no_dh, [2, 3, 4])] += share * C[no_dh][:, [5]]
    C[no_dh, 5] = 0
    # Krajevni odmik mešanice energentov: za vsako občino in vrsto stavbe razmerje med mešanico v izkaznicah občine in
    # državno mešanico v izkaznicah (površinsko uteženo, omiljeno z EI_SHRINK navideznimi stavbami državnega povprečja).
    # Tako hiše v mestih s plinovodom dobijo manj biomase kot hiše na podeželju; državne vsote uskladi korak spodaj.
    ei_cols = [f'ei_sh_{c}' for c in CARRIERS]
    for typ in ('HISA', 'BLOKI'):
        m_typ = (res.cat == typ).to_numpy()
        e = res.loc[m_typ & res.ei_sh_el.notna().to_numpy(), ['obcina', 'm2'] + ei_cols]
        nat = (e[ei_cols].mul(e.m2, axis=0).sum() / e.m2.sum()).to_numpy() + 1e-3
        g = e.groupby('obcina')
        loc = g.apply(lambda x: pd.Series((x[ei_cols].mul(x.m2, axis=0).sum() / x.m2.sum()).to_numpy())).to_numpy()
        cnt = g.size().to_numpy()[:, None]
        mix = (cnt * loc + EI_SHRINK * (nat - 1e-3)) / (cnt + EI_SHRINK) + 1e-3
        ratio = pd.DataFrame(mix / nat, index=g.size().index)
        rows = m_typ & np.isin(obc, ratio.index)
        C[rows] *= ratio.loc[obc[rows]].to_numpy()
    C *= (tot / np.maximum(C.sum(1), 1e-9))[:, None]
    ei_sh = res[[f'ei_sh_{c}' for c in CARRIERS]].to_numpy()
    has_ei = ~np.isnan(ei_sh).any(1)
    C[has_ei] = ei_sh[has_ei] * tot[has_ei, None]
    fuel = [2, 3, 4, 5]
    tc = res.es_tc.notna().to_numpy() & ~has_ei
    fsum = C[tc][:, fuel].sum(1)
    C[np.ix_(tc, fuel)] = 0
    C[tc, 0] += 0.3 * fsum; C[tc, 1] += 0.7 * fsum
    bio = res.es_biomasa.notna().to_numpy() & ~has_ei & ~tc
    fsum = C[bio][:, [2, 3, 5]].sum(1)
    C[np.ix_(bio, [2, 3, 5])] = 0
    C[bio, 4] += fsum
    # Začetni delež daljinske toplote 3 % pri vseh stavbah v občinah z omrežjem (tudi hiše so lahko priključene).
    seed = ~no_dh & (C[:, 5] < 0.03 * tot)
    C[seed] *= 0.97
    C[seed, 5] += 0.03 * tot[seed]
    # Dokazi (izkaznica, Eko sklad) določajo začetno razdelitev energentov; uskladitev na preglednico 30 velja za vse stavbe.
    fixed_c = np.zeros(n, bool)
    ctx.check(len(dh_m) >= 10, f'model: daljinska toplota v {len(dh_m)} občinah (po izkaznicah)')
    t30 = ctx.draft.table(r'Struktura končne rabe energije po energentih po scenariju NPS 2050')
    r30 = lambda p: num(row(t30, p)[2])
    target_c = np.array([r30(r'^Električna'), r30(r'^Toplota okolice') + r30(r'^Sončna'), r30(r'^Zemeljski plin') + r30(r'^Utekočinjeni') + r30(r'^Bioplin'),
                         r30(r'^Ekstra lahko'), r30(r'^Lesna'), r30(r'^Daljinska')]) * KWH_PER_KTOE
    target_c *= tot.sum() / target_c.sum()  # vsota ostane iz RenRates (bilanca); razlika zaokrožitev
    Cn = C / np.maximum(tot[:, None], 1e-9)
    Cn = _rake_rows(Cn, tot, target_c, fixed_c)
    C = Cn * tot[:, None]
    PE = C @ np.array([FP[c] for c in CARRIERS])

    # 6: razredi NPS in prag 43 % prek preslikave iz izkaznic
    nps_share = np.zeros((n, 7)); above = np.zeros(n)
    for typ in ('HISA', 'BLOKI'):
        c = next(x for x in fond['categories'] if x['id'] == typ)
        edges = [-np.inf] + [c['class_bounds'][b] for b in BOUNDS] + [np.inf]
        thr = c['worst_43']['threshold']
        ek = df.loc[(df.cat == typ) & df.ei_razred.notna() & df.ei_pe.notna()]
        M = np.ones((9, 7)); A = np.zeros(9)
        for i, h in enumerate(HC):
            v = ek.loc[ek.ei_razred == h, 'ei_pe']
            M[i] += np.bincount(pd.cut(v, edges, labels=False), minlength=7)
            A[i] = (v > thr).mean() if len(v) else 0.5
        M /= M.sum(1, keepdims=True)
        m_typ = (res.cat == typ).to_numpy()
        nps_share[m_typ] = P[m_typ] @ M
        above[m_typ] = P[m_typ] @ A

    out = pd.DataFrame({'obcina': obc, 'm2': res.m2.to_numpy(), 'fe': tot, 'pe': PE, 'above': above * res.m2.to_numpy(), 'known': known})
    for i, c in enumerate(CARRIERS):
        out[f'c_{c}'] = C[:, i]
    for i, k in enumerate(NPS):
        out[f'n_{k}'] = nps_share[:, i] * res.m2.to_numpy()
    return {'buildings': out, 'factor': factor0, 'target_c': target_c, 'dh_m': dh_m, 'Pe': Pe, 'arch': arch, 'w': w, 'cat': res.cat.to_numpy()}


def build(ctx: Context) -> dict:
    fond = json.loads((OUT / 'stavbni_fond.json').read_text(encoding='utf-8'))
    df = build_table(ctx)
    rr = load_renrates(ctx)
    tk = heat_map_by_municipality(ctx)
    runs = {k: run(ctx, df, rr, fond, tk, clip) for k, clip in VARIANTS.items()}
    c = runs['osrednja']['buildings']
    meas = build_measured(ctx, df)
    nres = model_nres.run(ctx, df, meas, rr, runs['osrednja']['dh_m'])

    # Kontrole
    fe_ktoe = c.fe.sum() / KWH_PER_KTOE
    t14 = ctx.draft.table(r'Končna in primarna raba energije po segmentih stavb v izhodiščnih letih')
    fe_ref, pe_ref = num(row(t14, r'^Stanovanjske stavbe')[2]), num(row(t14, r'^Stanovanjske stavbe')[4])
    ctx.check(abs(fe_ktoe - fe_ref) / fe_ref < 0.005, f'model: končna raba stanovanjskih stavb {fe_ktoe:.1f} ktoe ≈ osnutek {fe_ref} (±0,5 %)')
    pe_ktoe = c.pe.sum() / KWH_PER_KTOE
    ctx.check(abs(pe_ktoe - pe_ref) / pe_ref < 0.01, f'model: primarna raba stanovanjskih stavb {pe_ktoe:.1f} ktoe ≈ osnutek {pe_ref} (±1 %)')
    for i, cc in enumerate(CARRIERS):
        got = c[f'c_{cc}'].sum()
        ctx.check(abs(got / runs['osrednja']['target_c'][i] - 1) < 0.01, f'model: energent {cc} {got / KWH_PER_KTOE:.1f} ≈ preglednica 30 {runs["osrednja"]["target_c"][i] / KWH_PER_KTOE:.1f} ktoe (±1 %)')
    for k, r in runs.items():
        ctx.check(abs(r['buildings'].fe.sum() / c.fe.sum() - 1) < 1e-6, f'model: različica {k} ohrani državno vsoto')
    area = c.m2.sum()
    g_share = c.n_G.sum() / area
    above43 = c.above.sum() / area
    ctx.warn_unless(0.10 <= g_share <= 0.22, f'model: delež površine v razredu G NPS {100 * g_share:.1f} % (definicija: 15 % izkaznic)')
    ctx.warn_unless(0.35 <= above43 <= 0.50, f'model: delež površine nad pragom 43 % {100 * above43:.1f} % (definicija: 43 %)')

    def agg(g: pd.DataFrame, lo: pd.DataFrame, hi: pd.DataFrame) -> dict:
        a = g.m2.sum()
        fe = g.fe.sum()
        return {k: (float(v) if isinstance(v, np.floating) else v) for k, v in {
            'fe_gwh': round(fe / 1e6, 1), 'fe_gwh_lo': round(min(lo.fe.sum(), hi.fe.sum(), fe) / 1e6, 1), 'fe_gwh_hi': round(max(lo.fe.sum(), hi.fe.sum(), fe) / 1e6, 1),
            'fe_kwh_m2': round(fe / a), 'pe_gwh': round(g.pe.sum() / 1e6, 1), 'pe_kwh_m2': round(g.pe.sum() / a),
            'carriers_pct': {cc: float(round(100 * g[f'c_{cc}'].sum() / fe, 1)) for cc in CARRIERS},
            'nps_pct': {k: float(round(100 * g[f'n_{k}'].sum() / a, 1)) for k in NPS},
            'above43_area_pct': round(100 * g.above.sum() / a, 1),
            'known_pct': round(100 * g.known.mean(), 1),
        }.items()}

    ob = json.loads(ctx.files['obcine'].read_text(encoding='utf-8'))
    muni = {str(f['properties']['EID_OBCINA']): (int(f['properties']['SIFRA']), f['properties']['NAZIV']) for f in ob['features']}
    lo, hi = runs['brez']['buildings'], runs['mocna']['buildings']
    gl, gh = lo.groupby('obcina'), hi.groupby('obcina')
    out = {}
    for eid, g in c.groupby('obcina'):
        s, name = muni[eid]
        out[str(s)] = {'name': name, 'tk_factor': round(float(runs['osrednja']['factor'].get(eid, 1.0)), 3), **agg(g, gl.get_group(eid), gh.get_group(eid))}
    kz = kazalniki_obcin.build(ctx, c, nres, muni, DATA_DIR / 'raw' / 'prebivalci.json')
    for sfx, v in kz['municipalities'].items():
        out[sfx]['vse'] = v
    ixj = json.loads((OUT / 'obcine_index.json').read_text(encoding='utf-8'))
    ci = cilji_obcin.build(ctx, df, c, nres, muni, {'si': ixj['si'], 'm': {r['sifra']: json.loads((OUT / 'obcine' / f"{r['sifra']}.json").read_text(encoding='utf-8')) for r in ixj['municipalities']}})
    for sfx, v in ci.items():
        out[sfx]['cilji'] = v
    pj = projekcija_obcin.build(ctx, rr, runs['osrednja'], nres, {'sdh': kz['_sdh'], 's_bar': kz['_s_bar']}, muni)
    fos = (pd.concat([c[['obcina', 'c_gas', 'c_elko', 'fe']], nres[['obcina', 'c_gas', 'c_elko', 'fe']]]).groupby('obcina').sum())
    for sfx, v in pj.items():
        vse = out[sfx]['vse']
        eid = next(e for e, (s_, _) in muni.items() if str(s_) == sfx)
        f = fos.loc[eid]
        # izhodišče 2023 enako kot v bloku »vse« (ista umeritev), fosilna goriva iz izhodiščnih energentov
        v['2023'] = {'fe_gwh': vse['fe_gwh']['total'], 'fe_res_gwh': vse['fe_gwh']['res'], 'tgp_kt': vse['tgp_kt']['total'], 'ove_pct': vse['ove_pct']['total'],
                     'fossil_pct': round(100 * (f.c_gas + f.c_elko) / f.fe, 1)}
        out[sfx]['pot'] = v
    si_pot = {y: {'fe_gwh': round(sum(v[y]['fe_gwh'] for v in pj.values()), 1), 'tgp_kt': round(sum(v[y]['tgp_kt'] for v in pj.values()), 1)} for y in ('2023', '2030', '2040', '2050')}
    data = {
        'meta': ctx.meta([
            'Model RenRates (arhetipi, razporeditev po razredih toplotnih potreb, specifična končna energija, kalibracija λ na bilanco 2024)',
            'Kataster nepremičnin GURS 2026, register energetskih izkaznic 28. 9. 2026, Eko sklad do 2025',
            'Toplotna karta potrebne toplote za ogrevanje 2020 (mreža 100 m, vse stavbe)',
            f'{ctx.draft.name}: preglednice končne in primarne rabe po segmentih ter po energentih, emisij in OVE (2023)',
            'Energetsko knjigovodstvo javnega sektorja (poročila za leto 2025) in merjene energetske izkaznice',
            'AGEN-RS: Energetsko učinkoviti distribucijski sistemi toplote in hladu v letu 2025 (delež OVE po sistemih)',
            'SURS: prebivalstvo po občinah, 1. 1. 2026'],
            note='Ocena modela, ne meritev. Stanovanjske stavbe. Končna in primarna raba sta usklajeni z bilanco (država), po občinah pa razdeljeni '
                 'po stavbah glede na arhetip, izkaznice, podprte ukrepe, vpisane obnove in toplotno karto. Razpon: različice uskladitve s toplotno karto. '
                 'Razredi NPS in delež nad pragom 43 % so ocenjeni iz računskih razredov izkaznic (površina). '
                 'Blok »vse«: stanovanjske in nestanovanjske stavbe (javne in zasebne storitvene), emisije TGP in delež OVE, umerjeni na osnutek; '
                 'nestanovanjske stavbe z izmerjeno rabo iz energetskega knjigovodstva ali merjenih izkaznic, druge z modelom.'),
        'carriers': [{'id': 'el', 'name': 'električna energija'}, {'id': 'amb', 'name': 'toplota okolice in sončna energija'}, {'id': 'gas', 'name': 'plin (zemeljski, UNP)'},
                     {'id': 'elko', 'name': 'kurilno olje'}, {'id': 'bio', 'name': 'lesna biomasa'}, {'id': 'dh', 'name': 'daljinska toplota'}],
        'classes': NPS,
        'si': {**agg(c, lo, hi), 'vse': kz['si'], 'pot': si_pot},
        'pop_period': kz['pop_period'],
        'municipalities': out,
    }
    skupine_obcin.build(ctx, df)
    write_json('obcine_model', data)
    # Za zemljevid: modelska kazalnika dodamo v kazalo občin.
    ix = json.loads((OUT / 'obcine_index.json').read_text(encoding='utf-8'))
    def idx_fields(v):
        a = v['vse']
        return {'model_fe_kwh_m2': v['fe_kwh_m2'], 'model_above43_pct': v['above43_area_pct'],
                'k_fe_mwh_preb': a['fe_mwh_preb']['total'], 'k_res_fe_mwh_preb': a['fe_mwh_preb']['res'], 'k_tgp_t_preb': a['tgp_t_preb']['total'],
                'k_ove_pct': a['ove_pct']['total'], 'k_nres_kwh_m2': a['nres_kwh_m2'], 'pop': a['pop']}
    for r in ix['municipalities']:
        r.update(idx_fields(out[str(r['sifra'])]))
    ix['si'].update(idx_fields(data['si']))
    write_json('obcine_index', ix)
    write_csv('obcine_kazalniki', ['občina', 'šifra', 'prebivalci', 'končna energija vseh stavb [GWh]', 'stanovanjske [GWh]', 'nestanovanjske [GWh]',
                                   'končna energija na prebivalca [MWh]', 'stanovanjske na prebivalca [MWh]', 'emisije TGP [kt CO2 ekv]', 'emisije na prebivalca [t]',
                                   'delež OVE [%]', 'nestanovanjske [kWh/m²]', 'nestanovanjska površina z izmerjeno rabo [%]'],
              [[v['name'], s, v['vse']['pop'], v['vse']['fe_gwh']['total'], v['vse']['fe_gwh']['res'], v['vse']['fe_gwh']['nres'], v['vse']['fe_mwh_preb']['total'],
                v['vse']['fe_mwh_preb']['res'], v['vse']['tgp_kt']['total'], v['vse']['tgp_t_preb']['total'], v['vse']['ove_pct']['total'], v['vse']['nres_kwh_m2'],
                v['vse']['nres_measured_area_pct']] for s, v in sorted(out.items(), key=lambda x: x[1]['name'])])
    write_csv('obcine_model', ['občina', 'šifra', 'končna raba [GWh]', 'razpon od', 'razpon do', 'končna raba [kWh/m²]', 'primarna raba [GWh]',
                               'primarna raba [kWh/m²]', 'površina nad pragom 43 % [%]'] + [f'energent {cc} [%]' for cc in CARRIERS] + [f'razred {k} [% površine]' for k in NPS],
              [[v['name'], s, v['fe_gwh'], v['fe_gwh_lo'], v['fe_gwh_hi'], v['fe_kwh_m2'], v['pe_gwh'], v['pe_kwh_m2'], v['above43_area_pct']]
               + [v['carriers_pct'][cc] for cc in CARRIERS] + [v['nps_pct'][k] for k in NPS] for s, v in sorted(out.items(), key=lambda x: x[1]['name'])])
    return data
